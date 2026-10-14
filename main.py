import os
import sqlite3
from datetime import datetime

import flet as ft
from openpyxl import Workbook, load_workbook
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont

try:
    import arabic_reshaper
    from bidi.algorithm import get_display
except Exception:
    arabic_reshaper = None
    get_display = None


APP_TITLE = "مدرسة الهدى الثانوية بنين"
SUBJECTS = ["اللغة العربية", "الرياضيات", "اللغة الإنجليزية", "الفيزياء"]
DB_NAME = "students.db"


def get_data_dir(page: ft.Page):
    path = os.environ.get("FLET_APP_STORAGE_DATA")
    if path:
        try:
            os.makedirs(path, exist_ok=True)
            return path
        except Exception:
            pass
    path = os.path.join(os.path.expanduser("~"), ".huda_school")
    os.makedirs(path, exist_ok=True)
    return path


def get_output_dir(data_dir, is_android):
    """مجلد حفظ ملفات PDF وExcel."""
    candidates = []
    if is_android:
        candidates.append("/storage/emulated/0/Download")
    else:
        candidates.append(os.path.join(os.path.expanduser("~"), "Downloads"))
    for folder in candidates:
        try:
            if os.path.isdir(folder):
                test = os.path.join(folder, ".huda_test")
                with open(test, "w") as f:
                    f.write("x")
                os.remove(test)
                return folder
        except Exception:
            continue
    folder = os.path.join(data_dir, "exports")
    os.makedirs(folder, exist_ok=True)
    return folder


DB_PATH = None


def connect_db():
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    conn.execute("""
        CREATE TABLE IF NOT EXISTS students (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            name TEXT NOT NULL,
            class_name TEXT DEFAULT '',
            arabic REAL DEFAULT 0,
            math REAL DEFAULT 0,
            english REAL DEFAULT 0,
            physics REAL DEFAULT 0
        )
    """)
    conn.commit()
    return conn


def shape_arabic(text):
    text = str(text or "")
    if arabic_reshaper and get_display:
        try:
            return get_display(arabic_reshaper.reshape(text))
        except Exception:
            return text
    return text


def main(page: ft.Page):
    global DB_PATH
    data_dir = get_data_dir(page)
    DB_PATH = os.path.join(data_dir, DB_NAME)
    page.title = APP_TITLE
    page.rtl = True
    page.theme_mode = ft.ThemeMode.LIGHT
    page.padding = 14
    page.scroll = ft.ScrollMode.AUTO
    page.window_width = 1000
    page.window_height = 800

    is_android = page.platform == ft.PagePlatform.ANDROID
    out_dir = get_output_dir(data_dir, is_android)

    conn = connect_db()
    selected_student_id = None
    status = ft.Text("", color="#146C43")
    search_box = ft.TextField(label="بحث باسم الطالب أو الفصل", prefix_icon=ft.Icons.SEARCH, expand=True)
    name_box = ft.TextField(label="اسم الطالب", expand=True)
    class_box = ft.TextField(label="الفصل", width=150)
    arabic_box = ft.TextField(label="العربية", value="0", width=110, keyboard_type=ft.KeyboardType.NUMBER)
    math_box = ft.TextField(label="الرياضيات", value="0", width=110, keyboard_type=ft.KeyboardType.NUMBER)
    english_box = ft.TextField(label="الإنجليزية", value="0", width=110, keyboard_type=ft.KeyboardType.NUMBER)
    physics_box = ft.TextField(label="الفيزياء", value="0", width=110, keyboard_type=ft.KeyboardType.NUMBER)
    students_column = ft.Column(spacing=8)

    def show_message(message, error=False):
        status.value = message
        status.color = "#B42318" if error else "#146C43"
        try:
            page.open(ft.SnackBar(ft.Text(message)))
        except Exception:
            pass
        page.update()

    def number(value):
        try:
            return float(value or 0)
        except (ValueError, TypeError):
            return 0.0

    def clear_form():
        nonlocal selected_student_id
        selected_student_id = None
        name_box.value = ""
        class_box.value = ""
        arabic_box.value = "0"
        math_box.value = "0"
        english_box.value = "0"
        physics_box.value = "0"
        save_button.text = "إضافة الطالب"
        page.update()

    def refresh_students(*_):
        students_column.controls.clear()
        term = (search_box.value or "").strip()
        if term:
            rows = conn.execute(
                "SELECT * FROM students WHERE name LIKE ? OR class_name LIKE ? ORDER BY id DESC",
                (f"%{term}%", f"%{term}%")
            ).fetchall()
        else:
            rows = conn.execute("SELECT * FROM students ORDER BY id DESC").fetchall()

        if not rows:
            students_column.controls.append(
                ft.Container(
                    content=ft.Text("لا توجد بيانات بعد. أضف طالباً جديداً.", color="#667085"),
                    padding=20,
                    alignment=ft.alignment.center,
                )
            )
        for row in rows:
            total = sum(float(row[k] or 0) for k in ("arabic", "math", "english", "physics"))
            student_line = ft.Row(
                controls=[
                    ft.Text(str(row["id"]), width=40),
                    ft.Column([
                        ft.Text(row["name"], weight=ft.FontWeight.BOLD),
                        ft.Text(f"الفصل: {row['class_name'] or '—'}  |  المجموع: {total:g} / 400",
                                size=12, color="#475467"),
                    ], expand=True, spacing=3),
                    ft.IconButton(ft.Icons.EDIT, tooltip="تعديل", on_click=lambda e, r=dict(row): load_student(r)),
                    ft.IconButton(ft.Icons.PICTURE_AS_PDF, tooltip="شهادة PDF", on_click=lambda e, r=dict(row): make_pdf(r)),
                    ft.IconButton(ft.Icons.DELETE_OUTLINE, tooltip="حذف", icon_color="#B42318",
                                  on_click=lambda e, sid=row["id"], sn=row["name"]: confirm_delete(sid, sn)),
                ],
                vertical_alignment=ft.CrossAxisAlignment.CENTER,
            )
            students_column.controls.append(
                ft.Container(content=student_line, padding=10, border=ft.border.all(1, "#EAECF0"),
                             border_radius=10, bgcolor="#FFFFFF")
            )
        page.update()

    def load_student(row):
        nonlocal selected_student_id
        selected_student_id = int(row["id"])
        name_box.value = row["name"] or ""
        class_box.value = row["class_name"] or ""
        arabic_box.value = str(row["arabic"] or 0)
        math_box.value = str(row["math"] or 0)
        english_box.value = str(row["english"] or 0)
        physics_box.value = str(row["physics"] or 0)
        save_button.text = "حفظ التعديلات"
        page.update()

    def confirm_delete(student_id, student_name):
        def close_dialog(_):
            page.close(dialog)

        def delete_now(_):
            conn.execute("DELETE FROM students WHERE id=?", (student_id,))
            conn.commit()
            page.close(dialog)
            if selected_student_id == student_id:
                clear_form()
            refresh_students()
            show_message(f"تم حذف الطالب: {student_name}")

        dialog = ft.AlertDialog(
            modal=True,
            title=ft.Text("تأكيد الحذف"),
            content=ft.Text(f"هل تريد حذف الطالب «{student_name}»؟"),
            actions=[
                ft.TextButton("إلغاء", on_click=close_dialog),
                ft.FilledButton("حذف", on_click=delete_now, style=ft.ButtonStyle(bgcolor="#B42318")),
            ],
        )
        page.open(dialog)

    def save_student(_):
        nonlocal selected_student_id
        name = (name_box.value or "").strip()
        if not name:
            show_message("اكتب اسم الطالب أولاً.", True)
            return
        values = (
            name, (class_box.value or "").strip(),
            number(arabic_box.value), number(math_box.value),
            number(english_box.value), number(physics_box.value)
        )
        if selected_student_id is None:
            conn.execute(
                "INSERT INTO students(name,class_name,arabic,math,english,physics) VALUES(?,?,?,?,?,?)",
                values
            )
            message = "تمت إضافة الطالب وحفظ بياناته على الجهاز."
        else:
            conn.execute(
                "UPDATE students SET name=?,class_name=?,arabic=?,math=?,english=?,physics=? WHERE id=?",
                values + (selected_student_id,)
            )
            message = "تم حفظ التعديلات."
        conn.commit()
        clear_form()
        refresh_students()
        show_message(message)

    def make_pdf(row):
        try:
            path = os.path.join(out_dir, f"huda_student_{row['id']}.pdf")
            pdf = canvas.Canvas(path, pagesize=A4)
            width, height = A4

            base = os.path.dirname(os.path.abspath(__file__))
            candidates = [
                os.path.join(base, "Amiri-Regular.ttf"),
                os.path.join(base, "assets", "Amiri-Regular.ttf"),
            ]
            if os.environ.get("FLET_ASSETS_DIR"):
                candidates.append(os.path.join(os.environ["FLET_ASSETS_DIR"], "Amiri-Regular.ttf"))
            font_name = "Helvetica"
            for font_path in candidates:
                try:
                    if os.path.exists(font_path):
                        pdfmetrics.registerFont(TTFont("Amiri", font_path))
                        font_name = "Amiri"
                        break
                except Exception:
                    continue

            pdf.setFont(font_name, 18)
            pdf.drawCentredString(width / 2, height - 70, shape_arabic(APP_TITLE))
            pdf.setFont(font_name, 15)
            pdf.drawCentredString(width / 2, height - 110, shape_arabic("شهادة درجات طالب"))
            pdf.setFont(font_name, 12)
            y = height - 165
            lines = [
                f"اسم الطالب: {row['name']}",
                f"الفصل: {row['class_name'] or '—'}",
                f"اللغة العربية: {row['arabic'] or 0}",
                f"الرياضيات: {row['math'] or 0}",
                f"اللغة الإنجليزية: {row['english'] or 0}",
                f"الفيزياء: {row['physics'] or 0}",
                f"المجموع: {sum(float(row[k] or 0) for k in ('arabic','math','english','physics')):g} من 400",
                f"تاريخ الإصدار: {datetime.now().strftime('%Y-%m-%d')}",
            ]
            for line in lines:
                pdf.drawRightString(width - 55, y, shape_arabic(line))
                y -= 32
            pdf.line(55, y - 15, width - 55, y - 15)
            pdf.setFont(font_name, 10)
            pdf.drawCentredString(width / 2, 65, shape_arabic("إدارة المدرسة"))
            pdf.save()
        except Exception as ex:
            show_message(f"تعذر إنشاء PDF: {ex}", True)
            return
        show_message(f"تم إنشاء ملف PDF في: {path}")
        if not is_android:
            try:
                page.launch_url("file:///" + path.replace("\\", "/"))
            except Exception:
                pass

    def export_excel(_):
        try:
            rows = conn.execute(
                "SELECT id,name,class_name,arabic,math,english,physics FROM students ORDER BY id"
            ).fetchall()
            wb = Workbook()
            ws = wb.active
            ws.title = "الطلاب"
            ws.append(["رقم", "اسم الطالب", "الفصل", "اللغة العربية", "الرياضيات", "اللغة الإنجليزية", "الفيزياء"])
            for r in rows:
                ws.append([r["id"], r["name"], r["class_name"], r["arabic"], r["math"], r["english"], r["physics"]])
            path = os.path.join(out_dir, "huda_students.xlsx")
            wb.save(path)
            show_message(f"تم تصدير Excel إلى: {path}")
        except Exception as ex:
            show_message(f"تعذر التصدير: {ex}", True)

    def import_excel_result(e: ft.FilePickerResultEvent):
        if not e.files:
            return
        path = e.files[0].path
        if not path:
            show_message("تعذر الوصول إلى الملف المختار.", True)
            return
        if not path.lower().endswith(".xlsx"):
            show_message("اختر ملف Excel بامتداد .xlsx فقط.", True)
            return
        try:
            wb = load_workbook(path, data_only=True)
            ws = wb.active
            imported = 0
            for row in ws.iter_rows(min_row=2, values_only=True):
                if not row or len(row) < 2 or not row[1]:
                    continue
                name = str(row[1]).strip()
                class_name = str(row[2] or "") if len(row) > 2 else ""
                vals = [number(row[i]) if len(row) > i else 0 for i in range(3, 7)]
                conn.execute(
                    "INSERT INTO students(name,class_name,arabic,math,english,physics) VALUES(?,?,?,?,?,?)",
                    (name, class_name, *vals)
                )
                imported += 1
            conn.commit()
            refresh_students()
            show_message(f"تم استيراد {imported} طالباً.")
        except Exception as ex:
            show_message(f"تعذر استيراد الملف: {ex}", True)

    # FilePicker (flet 0.28.3): يُضاف إلى overlay ثم update
    file_picker = ft.FilePicker(on_result=import_excel_result)
    page.overlay.append(file_picker)
    page.update()

    def open_file_picker(e):
        try:
            file_picker.pick_files(
                allowed_extensions=["xlsx"],
                allow_multiple=False,
                file_type=ft.FilePickerFileType.CUSTOM,
            )
        except Exception as ex:
            show_message(f"تعذر فتح اختيار الملفات: {ex}", True)

    save_button = ft.FilledButton("إضافة الطالب", icon=ft.Icons.SAVE, on_click=save_student)

    header = ft.Container(
        content=ft.Column([
            ft.Text("بسم الله الرحمن الرحيم", size=16, text_align=ft.TextAlign.CENTER),
            ft.Text(APP_TITLE, size=25, weight=ft.FontWeight.BOLD, color="#FFFFFF",
                    text_align=ft.TextAlign.CENTER),
            ft.Text("نظام إدارة الطلاب والدرجات — يعمل دون إنترنت", size=13,
                    color="#F3E8FF", text_align=ft.TextAlign.CENTER),
        ], horizontal_alignment=ft.CrossAxisAlignment.CENTER, spacing=8),
        bgcolor="#6D28D9", padding=20, border_radius=14,
    )

    form = ft.Container(
        content=ft.Column([
            ft.Text("بيانات الطالب", size=18, weight=ft.FontWeight.BOLD, color="#5B21B6"),
            ft.Row([name_box, class_box], wrap=True),
            ft.Row([arabic_box, math_box, english_box, physics_box], wrap=True),
            ft.Row([
                save_button,
                ft.OutlinedButton("مسح الحقول", icon=ft.Icons.CLEAR, on_click=lambda e: clear_form()),
                ft.OutlinedButton("استيراد Excel", icon=ft.Icons.UPLOAD_FILE, on_click=open_file_picker),
                ft.OutlinedButton("تصدير Excel", icon=ft.Icons.DOWNLOAD, on_click=export_excel),
            ], wrap=True),
        ], spacing=12),
        padding=16, bgcolor="#F9F5FF", border_radius=12,
    )

    page.add(
        header,
        ft.Container(height=8),
        form,
        ft.Row([
            ft.Text("سجل الطلاب", size=20, weight=ft.FontWeight.BOLD),
            search_box,
        ], vertical_alignment=ft.CrossAxisAlignment.CENTER, wrap=True),
        status,
        students_column,
        ft.Text("تنبيه: البيانات تحفظ في قاعدة SQLite داخل مساحة التطبيق على هذا الجهاز.",
                size=11, color="#667085"),
    )
    search_box.on_change = refresh_students
    refresh_students()


if __name__ == "__main__":
    ft.app(target=main)
